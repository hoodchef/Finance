"""Builds the ASTS valuation workbook with live Excel formulas (Agent 1).

This is intentionally a separate implementation from engine.py; Agent 3's audit
recalculates this workbook in LibreOffice and diffs every scenario against the engine.
"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.comments import Comment
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.series import SeriesLabel

sys.path.insert(0, "/home/claude/asts"); sys.path.insert(0, "/home/claude/asts/data")
import inputs as I
import ibkr

# ------------------------------------------------------------------ styles
F = "Arial"
BLUE = Font(name=F, size=10, color="0000FF")
BLACK = Font(name=F, size=10, color="000000")
GREEN = Font(name=F, size=10, color="008000")
BOLD = Font(name=F, size=10, bold=True)
HDR = Font(name=F, size=10, bold=True, color="FFFFFF")
TITLE = Font(name=F, size=14, bold=True, color="1F3864")
SUB = Font(name=F, size=9, italic=True, color="595959")
GREY = Font(name=F, size=8, color="808080")
FILL_H = PatternFill("solid", fgColor="1F3864")
FILL_S = PatternFill("solid", fgColor="D9E1F2")
FILL_Y = PatternFill("solid", fgColor="FFFF00")
FILL_K = PatternFill("solid", fgColor="F2F2F2")
THIN = Side(style="thin", color="BFBFBF")
TOP = Border(top=Side(style="thin", color="000000"))
FMT = {"$mm": '#,##0.0;(#,##0.0);"-"', "%": '0.0%;(0.0%);"-"', "#": '#,##0;(#,##0);"-"',
       "x": '0.00x', "yr": '0', "yrs": '0.0', "mm": '#,##0.0;(#,##0.0);"-"', "$/mo": '$0.00',
       "flag": '0', "$": '$#,##0.00;($#,##0.00);"-"', "$mm/yr": '#,##0.0', "f": '0.0000',
       "sh": '#,##0;(#,##0);"-"', "pct2": '0.00%;(0.00%);"-"'}
SCOL = {0: "D", 1: "E", 2: "F", 3: "G"}
PCOLS = [CL(4 + i) for i in range(len(I.PERIODS))]      # D..R
TCOL = CL(4 + len(I.PERIODS))                           # S
wb = Workbook()


def setw(ws, widths):
    for k, v in widths.items():
        ws.column_dimensions[k].width = v


def title(ws, t, sub):
    ws["A1"] = t; ws["A1"].font = TITLE
    ws["A2"] = sub; ws["A2"].font = SUB
    ws.sheet_view.showGridLines = False


def section(ws, r, text, ncol=19):
    for c in range(1, ncol + 1):
        ws.cell(r, c).fill = FILL_S
    ws.cell(r, 1, text).font = BOLD


def put(ws, ref, val, font=BLACK, fmt=None, fill=None):
    c = ws[ref]; c.value = val; c.font = font
    if fmt: c.number_format = FMT.get(fmt, fmt)
    if fill: c.fill = fill
    return c


# ================================================================== INPUTS
wsI = wb.active; wsI.title = "Inputs"
title(wsI, "INPUTS & ASSUMPTIONS", "Blue = hard-coded input (edit here only). Yellow = key value driver. "
      "Every formula in the model references this sheet. Units in column C. $mm unless stated.")
setw(wsI, {"A": 13, "B": 52, "C": 8, "D": 11, "E": 11, "F": 11, "G": 11, "H": 95})
KEY_DRIVERS = {"take_t1", "arpu_t1", "prob", "unit_cost", "fl_ss", "cip_credit", "t1_start", "rf", "erp", "g",
               "sat_life", "take_t2"}
IN = {}   # key -> (row, is_scenario)
r = 4
put(wsI, "A4", "Key", BOLD); put(wsI, "B4", "Global parameter", BOLD); put(wsI, "C4", "Unit", BOLD)
put(wsI, "D4", "Value", BOLD); put(wsI, "H4", "Source / rationale", BOLD)
r = 5
for key, lab, unit, val, src in I.GLOBAL:
    if key == "hdr":
        section(wsI, r, lab, 8); r += 1; continue
    wsI.cell(r, 1, key).font = GREY
    wsI.cell(r, 2, lab).font = BLACK
    wsI.cell(r, 3, unit).font = GREY
    c = wsI.cell(r, 4, val); c.font = BLUE; c.number_format = FMT.get(unit, "General")
    if unit == "%" and abs(val) < 1 and key in ("rf", "erp", "kd_hi", "kd_term"):
        c.number_format = FMT["pct2"]
    if key in KEY_DRIVERS: c.fill = FILL_Y
    wsI.cell(r, 8, src).font = GREY
    IN[key] = (r, False); r += 1
r += 1
put(wsI, f"A{r}", "Key", BOLD); put(wsI, f"B{r}", "Scenario parameter", BOLD); put(wsI, f"C{r}", "Unit", BOLD)
for s in range(4):
    put(wsI, f"{SCOL[s]}{r}", I.SCEN[s], BOLD)
put(wsI, f"H{r}", "Source / rationale", BOLD)
SCEN_HDR_ROW = r
r += 1
for key, lab, unit, vals, src in I.SCENARIO:
    if key == "hdr":
        section(wsI, r, lab, 8); r += 1; continue
    wsI.cell(r, 1, key).font = GREY
    wsI.cell(r, 2, lab).font = BLACK
    wsI.cell(r, 3, unit).font = GREY
    for s in range(4):
        c = wsI.cell(r, 4 + s, vals[s]); c.font = BLUE; c.number_format = FMT.get(unit, "General")
        if key in KEY_DRIVERS: c.fill = FILL_Y
    wsI.cell(r, 8, src).font = GREY
    IN[key] = (r, True); r += 1
PROB_ROW = IN["prob"][0]
r += 1
put(wsI, f"B{r}", "Sum of scenario probabilities (must = 100%)", BOLD)
put(wsI, f"D{r}", f"=SUM(D{PROB_ROW}:G{PROB_ROW})", BLACK, "%")
PROB_SUM = f"Inputs!$D${r}"
wsI.freeze_panes = "D5"


def IREF(key, s=None):
    row, isscen = IN[key]
    col = SCOL[s] if isscen else "D"
    return f"Inputs!${col}${row}"


# ================================================================== MARKET DATA
wsM = wb.create_sheet("Market_Data")
title(wsM, "MARKET DATA (Interactive Brokers)", "Weekly closes, 2-yr window, retrieved from IBKR 10-Sep-2026. "
      "ASTS conid 480745767 (NASDAQ), SPY conid 756733 (ARCA). Last bar is a partial week (excluded from regression).")
setw(wsM, {"A": 12, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 3, "I": 44, "J": 13, "K": 60})
for j, h in enumerate(["Week start", "ASTS close", "SPY close", "ASTS ret", "SPY ret", "ASTS peak", "Drawdown"]):
    put(wsM, f"{CL(1 + j)}5", h, HDR, fill=FILL_H)
M0 = 6
n = len(ibkr.TIME)
for i in range(n):
    rr = M0 + i
    put(wsM, f"A{rr}", ibkr.TIME[i], BLACK)
    put(wsM, f"B{rr}", ibkr.ASTS[i], BLUE, "$")
    put(wsM, f"C{rr}", ibkr.SPY[i], BLUE, "$")
    if i > 0:
        put(wsM, f"D{rr}", f"=B{rr}/B{rr-1}-1", BLACK, "%")
        put(wsM, f"E{rr}", f"=C{rr}/C{rr-1}-1", BLACK, "%")
        put(wsM, f"F{rr}", f"=MAX(F{rr-1},B{rr})", BLACK, "$")
    else:
        put(wsM, f"F{rr}", f"=B{rr}", BLACK, "$")
    put(wsM, f"G{rr}", f"=B{rr}/F{rr}-1", BLACK, "%")
ML = M0 + n - 1                      # last row (partial week)
RA = f"D{M0+1}:D{ML-1}"; RS = f"E{M0+1}:E{ML-1}"
stats = [
    ("Current price (last close, $)", f"=B{ML}", "$", "price"),
    ("Full-week returns in regression (n)", f"=COUNT({RA})", "#", "n"),
    ("Raw beta vs SPY (weekly, 2y)", f"=SLOPE({RA},{RS})", "f", "beta_raw"),
    ("Std. error of beta", f"=STEYX({RA},{RS})/(STDEV({RS})*SQRT(COUNT({RS})-1))", "f", "beta_se"),
    ("t-statistic", "=J9/J10", "f", "beta_t"),
    ("R-squared (systematic share of variance)", f"=RSQ({RA},{RS})", "%", "r2"),
    ("Blume-adjusted beta (used in WACC)", f"={IREF('blume_w')}*J9+(1-{IREF('blume_w')})", "f", "beta_adj"),
    ("ASTS annualised volatility", f"=STDEV({RA})*SQRT(52)", "%", "vol"),
    ("SPY annualised volatility", f"=STDEV({RS})*SQRT(52)", "%", "spy_vol"),
    ("Idiosyncratic volatility", "=J14*SQRT(1-J12)", "%", "idio"),
    ("Max drawdown (2y)", f"=MIN(G{M0}:G{ML})", "%", "mdd"),
    ("1-week historical VaR 95% (return)", f"=PERCENTILE({RA},0.05)", "%", "var95"),
    ("1-week CVaR 95% (mean of returns <= VaR)", f"=AVERAGEIF({RA},\"<=\"&J18)", "%", "cvar95"),
    ("1-week parametric VaR 95% (normal)", f"=AVERAGE({RA})-1.645*STDEV({RA})", "%", "pvar95"),
    ("Annualised downside deviation (MAR 0)", f"=SQRT(SUMPRODUCT(({RA}<0)*({RA})^2)/COUNT({RA}))*SQRT(52)", "%", "dd"),
    ("2-yr high close", f"=MAX(B{M0}:B{ML})", "$", "hi"),
    ("2-yr low close", f"=MIN(B{M0}:B{ML})", "$", "lo"),
]
MS = {}
put(wsM, "I5", "Statistic", HDR, fill=FILL_H); put(wsM, "J5", "Value", HDR, fill=FILL_H)
put(wsM, "K5", "Note", HDR, fill=FILL_H)
for j, (lab, f, fmt, k) in enumerate(stats):
    rr = 7 + j
    put(wsM, f"I{rr}", lab); put(wsM, f"J{rr}", f, BLACK, fmt); MS[k] = f"Market_Data!$J${rr}"
wsM["K9"] = "95% CI roughly beta +/- 1.96 x SE: statistically weak (R-01)."; wsM["K9"].font = GREY
wsM["K13"] = "Blume (1971): 0.67 x raw + 0.33. Chosen over raw/shrinkage - see Review_Log R-01."; wsM["K13"].font = GREY
wsM.freeze_panes = "A6"

# ================================================================== CAPITAL STRUCTURE
wsC = wb.create_sheet("Capital_Structure")
title(wsC, "CAPITAL STRUCTURE & EQUITY BRIDGE INPUTS", "Balance sheet at 30-Jun-2026 (10-Q) pro forma for the Jul-2026 "
      "$1.15bn 2034 convertible. Share count at 6-Aug-2026. $mm; shares in mm unless stated.")
setw(wsC, {"A": 50, "B": 15, "C": 13, "D": 13, "E": 13, "F": 13, "G": 70})
r = 4; section(wsC, r, "Share count (economic interests)", 7); r += 1
for h, col in zip(["Class", "Shares"], ["A", "B"]):
    put(wsC, f"{col}{r}", h, BOLD)
put(wsC, f"G{r}", "Source", BOLD); r += 1
s_first = r
for lab, cnt, src in I.SHARES:
    put(wsC, f"A{r}", lab); put(wsC, f"B{r}", cnt, BLUE, "sh"); put(wsC, f"G{r}", src, GREY); r += 1
put(wsC, f"A{r}", "Basic economic shares (mm) - S0", BOLD)
put(wsC, f"B{r}", f"=SUM(B{s_first}:B{r-1})/1000000", BLACK, "mm"); S0REF = f"Capital_Structure!$B${r}"
put(wsC, f"G{r}", "Up-C: Class B/C are voting-only; economics sit in AST LLC units exchangeable 1:1 (R-05).", GREY)
r += 2; section(wsC, r, "Cash & liquid assets", 7); r += 1
put(wsC, f"A{r}", "Item", BOLD); put(wsC, f"B{r}", "$mm", BOLD); put(wsC, f"C{r}", "Include (1/0)", BOLD)
put(wsC, f"D{r}", "Included $mm", BOLD); r += 1
c_first = r
for lab, amt, inc, src in I.CASH:
    put(wsC, f"A{r}", lab); put(wsC, f"B{r}", amt, BLUE, "$mm"); put(wsC, f"C{r}", inc, BLUE, "flag")
    put(wsC, f"D{r}", f"=B{r}*C{r}", BLACK, "$mm"); put(wsC, f"G{r}", src, GREY); r += 1
put(wsC, f"A{r}", "Total cash & liquid assets", BOLD); put(wsC, f"D{r}", f"=SUM(D{c_first}:D{r-1})", BOLD, "$mm")
CASHREF = f"Capital_Structure!$D${r}"; CASH_ROW = r; r += 1
put(wsC, f"A{r}", "Check: pro forma cash+restricted vs company-reported 'over $3.7bn'")
put(wsC, f"D{r}", f"=D{c_first}+D{c_first+1}+B{c_first+2}+D{c_first+3}+D{c_first+4}+D{c_first+5}", BLACK, "$mm")
PF_CASH = f"Capital_Structure!$D${r}"
put(wsC, f"G{r}", "Company: 'over $3.7 billion ... pro forma for convertible notes offering' (Q2-26 deck).", GREY)
r += 2; section(wsC, r, "Non-convertible debt (face)", 7); r += 1
d_first = r
for lab, amt, src in I.DEBT:
    put(wsC, f"A{r}", lab); put(wsC, f"B{r}", amt, BLUE, "$mm"); put(wsC, f"G{r}", src, GREY); r += 1
put(wsC, f"A{r}", "Total non-convertible debt", BOLD); put(wsC, f"B{r}", f"=SUM(B{d_first}:B{r-1})", BOLD, "$mm")
DEBTREF = f"Capital_Structure!$B${r}"; r += 2
section(wsC, r, "Convertible notes", 7); r += 1
for h, col in zip(["Tranche", "Face $mm", "Conv. price $", "Cap price $", "Shares mm", "Trigger K $"],
                  ["A", "B", "C", "D", "E", "F"]):
    put(wsC, f"{col}{r}", h, BOLD)
put(wsC, f"G{r}", "Source", BOLD); r += 1
cv_first = r
for lab, face, cp, cap, src in I.CONVERTS:
    put(wsC, f"A{r}", lab); put(wsC, f"B{r}", face, BLUE, "$mm"); put(wsC, f"C{r}", cp, BLUE, "$")
    put(wsC, f"D{r}", cap, BLUE, "$"); put(wsC, f"E{r}", f"=B{r}/C{r}", BLACK, "f")
    put(wsC, f"F{r}", f"=IF(D{r}>0,D{r},C{r})", BLACK, "$"); put(wsC, f"G{r}", src, GREY); r += 1
cv_last = r - 1
put(wsC, f"A{r}", "Total convertible face", BOLD); put(wsC, f"B{r}", f"=SUM(B{cv_first}:B{cv_last})", BOLD, "$mm")
CONVREF = f"Capital_Structure!$B${r}"; r += 1
put(wsC, f"A{r}", "Check: total debt face at 30-Jun-2026 vs 10-Q ($3,022.2mm)")
put(wsC, f"B{r}", f"={DEBTREF}+{CONVREF}-B{cv_last}", BLACK, "$mm"); DEBT_CHK = f"Capital_Structure!$B${r}"
put(wsC, f"G{r}", "10-Q Note 6 total debt face $3,022,152k (excludes Jul-2026 notes).", GREY)
r += 2; section(wsC, r, "Equity awards (treasury-stock logic via value-consistent dilution engine)", 7); r += 1
for h, col in zip(["Award", "Count", "Strike $"], ["A", "B", "C"]):
    put(wsC, f"{col}{r}", h, BOLD)
put(wsC, f"G{r}", "Source", BOLD); r += 1
aw_first = r
for lab, cnt, k, src in I.AWARDS:
    put(wsC, f"A{r}", lab); put(wsC, f"B{r}", cnt, BLUE, "sh"); put(wsC, f"C{r}", k, BLUE, "$")
    put(wsC, f"G{r}", src, GREY); r += 1
aw_last = r - 1
r += 1; section(wsC, r, "Dilutive instrument table (feeds every Calc sheet)", 7); r += 1
for h, col in zip(["Instrument", "Trigger K $", "Shares n (mm)", "Value added on conversion D $mm"], ["A", "B", "C", "D"]):
    put(wsC, f"{col}{r}", h, BOLD)
put(wsC, f"G{r}", "Convert: D = face (plain) or n x cap (capped call). Option: D = n x strike (cash in). RSU: K=0, D=0.", GREY)
r += 1
INST_FIRST = r
for i in range(cv_first, cv_last + 1):
    put(wsC, f"A{r}", f"=A{i}", GREEN); put(wsC, f"B{r}", f"=F{i}", BLACK, "$")
    put(wsC, f"C{r}", f"=E{i}", BLACK, "f"); put(wsC, f"D{r}", f"=C{r}*B{r}", BLACK, "$mm"); r += 1
for i in range(aw_first, aw_last + 1):
    put(wsC, f"A{r}", f"=A{i}", GREEN); put(wsC, f"B{r}", f"=C{i}", BLACK, "$")
    put(wsC, f"C{r}", f"=B{i}/1000000", BLACK, "f"); put(wsC, f"D{r}", f"=C{r}*B{r}", BLACK, "$mm"); r += 1
INST_LAST = r - 1
N_INST = INST_LAST - INST_FIRST + 1
r += 1; section(wsC, r, "Market-implied reference points", 7); r += 1
put(wsC, f"A{r}", "Equity market value at current price ($mm, basic)")
put(wsC, f"B{r}", f"={MS['price']}*{S0REF}", BLACK, "$mm"); MKTCAP = f"Capital_Structure!$B${r}"; r += 1
put(wsC, f"A{r}", "Market-implied EV ($mm) = mkt cap + all debt face - cash")
put(wsC, f"B{r}", f"={MKTCAP}+{DEBTREF}+{CONVREF}-{CASHREF}", BLACK, "$mm"); MKT_EV = f"Capital_Structure!$B${r}"; r += 1
put(wsC, f"A{r}", "Debt / (debt + equity) at market")
put(wsC, f"B{r}", f"=({DEBTREF}+{CONVREF})/({DEBTREF}+{CONVREF}+{MKTCAP})", BLACK, "%"); DV_MKT = f"Capital_Structure!$B${r}"

# ================================================================== WACC
wsW = wb.create_sheet("WACC")
title(wsW, "DISCOUNT RATE", "CAPM cost of equity; beta estimated live from IBKR data (Market_Data). Growth-phase WACC fades "
      "linearly to a mature WACC. Same WACC for every scenario - scenario risk is in cash flows, not the rate (R-03).")
setw(wsW, {"A": 46, "B": 10, "C": 12})
for i in range(len(I.PERIODS)):
    wsW.column_dimensions[PCOLS[i]].width = 9
rows = [
    ("Risk-free rate", f"={IREF('rf')}", "pct2", "rf"),
    ("Equity risk premium", f"={IREF('erp')}", "pct2", "erp"),
    ("Beta - growth phase (Blume-adjusted, live)", f"={MS['beta_adj']}", "f", "b_hi"),
    ("Cost of equity - growth phase", "=C5+C7*C6", "pct2", "ke_hi"),
    ("Debt / capital - growth phase", f"={IREF('dv_hi')}", "%", "dv_hi"),
    ("After-tax cost of debt - growth phase", f"={IREF('kd_hi')}*(1-{IREF('t_hi')})", "pct2", "kd_hi"),
    ("WACC - growth phase", f"=(1-C9)*C8+C9*C10+{IREF('wacc_shift')}", "pct2", "w_hi"),
    ("Beta - mature phase", f"={IREF('beta_term')}", "f", "b_t"),
    ("Cost of equity - mature phase", "=C5+C12*C6", "pct2", "ke_t"),
    ("Debt / capital - mature phase", f"={IREF('dv_term')}", "%", "dv_t"),
    ("After-tax cost of debt - mature phase", f"={IREF('kd_term')}*(1-{IREF('t_term')})", "pct2", "kd_t"),
    ("WACC - mature / terminal", f"=(1-C14)*C13+C14*C15+{IREF('wacc_shift')}", "pct2", "w_t"),
]
WR = {}
for j, (lab, f, fmt, k) in enumerate(rows):
    rr = 5 + j
    put(wsW, f"A{rr}", lab); put(wsW, f"C{rr}", f, GREEN if "Inputs" in f or "Market" in f else BLACK, fmt)
    WR[k] = f"WACC!$C${rr}"
put(wsW, "A19", "Period", BOLD); put(wsW, "A20", "Year", BOLD); put(wsW, "A21", "WACC applied", BOLD)
for i, per in enumerate(I.PERIODS):
    c = PCOLS[i]
    put(wsW, f"{c}19", per, BOLD); put(wsW, f"{c}20", I.YEARS[i], BLACK, "yr")
    fs, fe = IREF("wacc_fade_start"), IREF("wacc_fade_end")
    put(wsW, f"{c}21", f"=IF({c}20<{fs},$C$11,IF({c}20>={fe},$C$16,$C$11+($C$16-$C$11)*({c}20-{fs}+1)/({fe}-{fs}+1)))",
        BLACK, "pct2")
WACC_ROW = 21

# ================================================================== CALC SHEETS
CALC_ROWS = {}


def build_calc(s):
    ws = wb.create_sheet(f"Calc_{I.SCEN[s]}")
    title(ws, f"MODEL - {I.SCEN[s].upper()} SCENARIO", "Live formulas. Inputs from 'Inputs' column "
          f"{SCOL[s]} (scenario) / D (global). $mm unless stated. Periods: 2026H2 stub + 2027-2040; col S = 2041 normalised.")
    setw(ws, {"A": 46, "B": 8, "C": 12})
    for c in PCOLS + [TCOL]:
        ws.column_dimensions[c].width = 10
    ws.freeze_panes = "D6"
    R = {}
    rowno = [4]

    def nxt():
        rowno[0] += 1
        return rowno[0]

    def X(key):
        return IREF(key, s)

    def row(key, label, unit, fgen, fmt=None, bold=False, first=None, scalar=None, font=None):
        rr = nxt(); R[key] = rr
        ws.cell(rr, 1, label).font = BOLD if bold else BLACK
        ws.cell(rr, 2, unit).font = GREY
        for i, c in enumerate(PCOLS):
            pc = PCOLS[i - 1] if i > 0 else None
            f = first(c) if (i == 0 and first is not None) else fgen(c, pc, i)
            if f is None:
                continue
            cell = ws[f"{c}{rr}"]; cell.value = f
            cell.font = font or (BOLD if bold else (GREEN if isinstance(f, str) and ("Inputs!" in f or "WACC!" in f)
                                                   and f.count("!") and not any(ch in f for ch in "+-*/") else BLACK))
            cell.number_format = FMT.get(fmt or unit, "General")
            if bold: cell.border = TOP
        if scalar is not None:
            cell = ws[f"C{rr}"]; cell.value = scalar; cell.font = BLACK
            cell.number_format = FMT.get(fmt or unit, "General")
        return rr

    def sec(text):
        rr = nxt(); section(ws, rr, text)

    y = lambda c: f"{c}${R['year']}"
    # --- timeline
    row("period", "Period", "", lambda c, pc, i: I.PERIODS[i], font=BOLD)
    ws[f"{TCOL}{R['period']}"] = "2041N"; ws[f"{TCOL}{R['period']}"].font = BOLD
    row("year", "Calendar year", "yr", lambda c, pc, i: I.YEARS[i])
    row("len", "Period length (years)", "yrs", lambda c, pc, i: I.LENS[i])
    row("start", "Period start (year fraction)", "yrs", lambda c, pc, i: I.STARTS[i])
    row("mid", "Cash-flow timing (mid-period)", "yrs", lambda c, pc, i: f"={c}{R['start']}+{c}{R['len']}/2")
    for k in ("year", "len", "start"):
        for c in PCOLS:
            ws[f"{c}{R[k]}"].font = BLUE
    # --- fleet
    sec("CONSTELLATION (Block 2 satellites)")
    life = f"ROUND({X('sat_life')},0)"
    row("tgt", "Target fleet, end of period", "#",
        lambda c, pc, i: f"=IF({y(c)}=2026,{X('fl_26')},IF({y(c)}=2027,{X('fl_27')},IF({y(c)}=2028,{X('fl_28')},"
                         f"IF({y(c)}=2029,{X('fl_29')},{X('fl_ss')}))))")
    rb = rowno[0] + 1
    row("beg", "Beginning fleet", "#", lambda c, pc, i: f"={pc}{rb+3}", first=lambda c: f"={X('fleet_open')}")
    rN = rb + 2
    row("ret", "Retirements (end of design life)", "#",
        lambda c, pc, i: f"=IF({y(c)}-{life}=2026,$D${rN}+{X('fleet_open')},IF({y(c)}-{life}>=2027,"
                         f"INDEX($D${rN}:{pc}{rN},1,{y(c)}-{life}-2026+1),0))",
        first=lambda c: "=0")
    row("N", "Successful launches (net adds + replacements)", "#",
        lambda c, pc, i: f"=MAX(0,{c}{R['tgt']}-{c}{R['beg']}+{c}{R['ret']})")
    row("end", "Ending fleet", "#", lambda c, pc, i: f"={c}{R['beg']}-{c}{R['ret']}+{c}{R['N']}", bold=True)
    assert R["N"] == rN and R["end"] == rb + 3
    row("G", "Satellites built incl. launch losses", "#", lambda c, pc, i: f"={c}{R['N']}/(1-{X('loss_rate')})", fmt="$mm")
    row("avg", "Average fleet in orbit", "#", lambda c, pc, i: f"=({c}{R['beg']}+{c}{R['end']})/2", fmt="$mm")
    row("U", "Unit cost incl. launch ($mm/sat)", "$mm",
        lambda c, pc, i: f"={X('unit_cost')}*(1+{X('unit_infl')})^MAX(0,{y(c)}-2028)")
    row("SC", "Satellite capex (gross, placed in service)", "$mm", lambda c, pc, i: f"={c}{R['G']}*{c}{R['U']}")
    rr = nxt(); R["pool"] = rr
    ws.cell(rr, 1, "CIP pre-funding pool (scalar, col C)").font = BLACK
    ws[f"C{rr}"] = f"={X('cip_sat')}*{X('cip_credit')}"; ws[f"C{rr}"].number_format = FMT["$mm"]
    rC = rr + 1
    row("C", "CIP credit applied (already-paid capex)", "$mm",
        lambda c, pc, i: f"=MIN({c}{R['SC']},$C${R['pool']}-SUM($D${rC}:{pc}{rC}))",
        first=lambda c: f"=MIN({c}{R['SC']},$C${R['pool']})")
    row("netSC", "Satellite capex, cash", "$mm", lambda c, pc, i: f"={c}{R['SC']}-{c}{R['C']}", bold=True)
    # --- revenue
    sec("REVENUE BUILD")
    row("cov1", "Tier-1 coverage (fraction of partner base)", "%",
        lambda c, pc, i: f"=MIN(1,MAX(0,({c}{R['avg']}-{X('s0_t1')})/({X('sf_t1')}-{X('s0_t1')})))")
    row("cov2", "Tier-2 coverage", "%",
        lambda c, pc, i: f"=MIN(1,MAX(0,({c}{R['avg']}-{X('s0_t2')})/({X('sf_t2')}-{X('s0_t2')})))")
    rr = nxt(); R["k"] = rr
    ws.cell(rr, 1, "Logistic steepness k = 2 ln(19) / ramp (col C)").font = BLACK
    ws[f"C{rr}"] = f"=2*LN(19)/{X('ramp')}"; ws[f"C{rr}"].number_format = FMT["f"]
    for tier in ("1", "2"):
        st, pk = X(f"t{tier}_start"), X(f"take_t{tier}")
        row(f"ad{tier}", f"Tier-{tier} adoption (take rate of covered subs)", "%",
            lambda c, pc, i, st=st, pk=pk: f"=IF({c}{R['mid']}>={st},{pk}/(1+EXP(-$C${R['k']}*({c}{R['mid']}-({st}+{X('ramp')}/2)))),0)",
            fmt="pct2")
    row("base", "MNO partner subscriber base (mm)", "mm",
        lambda c, pc, i: f"={X('partner_subs')}*(1+{X('sub_g')})^MAX(0,{y(c)}-2026)")
    row("subs1", "Tier-1 paying users (mm, period avg)", "mm",
        lambda c, pc, i: f"={c}{R['base']}*{X('t1_share')}*{c}{R['cov1']}*{c}{R['ad1']}")
    row("subs2", "Tier-2 paying users (mm, period avg)", "mm",
        lambda c, pc, i: f"={c}{R['base']}*(1-{X('t1_share')})*{c}{R['cov2']}*{c}{R['ad2']}")
    row("px", "ARPU index (2027 = 1.00)", "x",
        lambda c, pc, i: f"=(1+{X('px_g')})^MIN(MAX(0,{y(c)}-2027),{X('px_end')}-2027)*(1+{X('arpu_lr_g')})^MAX(0,{y(c)}-{X('px_end')})",
        fmt="f")
    row("rev1", "Tier-1 service revenue", "$mm",
        lambda c, pc, i: f"={c}{R['subs1']}*{X('arpu_t1')}*{c}{R['px']}*12*{c}{R['len']}")
    row("rev2", "Tier-2 service revenue", "$mm",
        lambda c, pc, i: f"={c}{R['subs2']}*{X('arpu_t2')}*{c}{R['px']}*12*{c}{R['len']}")
    row("gov", "Government revenue", "$mm",
        lambda c, pc, i: f"=IF({y(c)}=2026,{X('gov_26')},IF({y(c)}=2027,{X('gov_27')},IF({y(c)}=2028,{X('gov_28')},"
                         f"IF({y(c)}=2029,{X('gov_29')},{X('gov_30')}*(1+{X('gov_g')})^({y(c)}-2030)))))")
    row("prod", "Gateway / product revenue", "$mm",
        lambda c, pc, i: f"=IF({y(c)}=2026,{X('prod_26')},IF({y(c)}=2027,{X('prod_27')},IF({y(c)}=2028,{X('prod_28')},{X('prod_ss')})))")
    row("rev", "Total revenue", "$mm", lambda c, pc, i: f"=SUM({c}{R['rev1']}:{c}{R['prod']})", bold=True)
    row("growth", "Revenue growth (annualised)", "%",
        lambda c, pc, i: f"=IF({pc}{R['rev']}=0,0,({c}{R['rev']}/{c}{R['len']})/({pc}{R['rev']}/{pc}{R['len']})-1)",
        first=lambda c: f"=IF({X('h1_rev')}=0,0,({c}{R['rev']}/{c}{R['len']})/({X('h1_rev')}/0.5)-1)")
    row("subspersat", "Paying users per satellite in orbit (mm)", "mm",
        lambda c, pc, i: f"=IF({c}{R['avg']}=0,0,({c}{R['subs1']}+{c}{R['subs2']})/{c}{R['avg']})", fmt="f")
    # --- costs
    sec("OPERATING COSTS & EBITDA")
    row("cogs", "Cost of revenue (gateway + government)", "$mm",
        lambda c, pc, i: f"={c}{R['prod']}*(1-{X('prod_margin')})+{c}{R['gov']}*{X('gov_cogs')}")
    row("var", "Service variable cost", "$mm", lambda c, pc, i: f"=({c}{R['rev1']}+{c}{R['rev2']})*{X('var_cost')}")
    row("lig", "Ligado revenue share", "$mm", lambda c, pc, i: f"={c}{R['rev1']}*{X('ligado_share')}")
    row("spec", "Spectrum usage payments", "$mm",
        lambda c, pc, i: f"=IF({y(c)}>={X('ligado_year')},{X('spec_fee')}*{c}{R['len']},0)")
    row("opex", "Fixed adj. opex (ex-COGS, ex-SBC)", "$mm",
        lambda c, pc, i: f"=IF({y(c)}=2027,{X('opex_27')},{pc}{R['opex']+0}*(1+IF({y(c)}=2028,{X('opex_g28')},"
                         f"IF({y(c)}=2029,{X('opex_g29')},IF({y(c)}=2030,{X('opex_g30')},{X('opex_gss')})))))",
        first=lambda c: f"={X('opex_26')}")
    row("netx", "Network, ground & insurance opex", "$mm",
        lambda c, pc, i: f"={c}{R['avg']}*{X('net_opex_sat')}*{c}{R['len']}")
    row("sbc", "Stock-based compensation (treated as cost)", "$mm",
        lambda c, pc, i: f"=MAX({X('sbc_27')}*(1-{X('sbc_decay')})^MAX(0,{y(c)}-2027),{X('sbc_floor')}*{c}{R['rev']})",
        first=lambda c: f"={X('sbc_26')}")
    row("costs", "Total operating costs", "$mm", lambda c, pc, i: f"=SUM({c}{R['cogs']}:{c}{R['sbc']})", bold=True)
    row("ebitda", "EBITDA (after SBC)", "$mm", lambda c, pc, i: f"={c}{R['rev']}-{c}{R['costs']}", bold=True)
    row("margin", "EBITDA margin", "%", lambda c, pc, i: f"=IF({c}{R['rev']}=0,0,{c}{R['ebitda']}/{c}{R['rev']})")
    # --- capex & other
    sec("CAPEX & OTHER INVESTING ITEMS")
    row("ocx", "Other capex (ground, facilities)", "$mm",
        lambda c, pc, i: f"=IF({y(c)}=2026,{X('ocx_26')},IF({y(c)}=2027,{X('ocx_27')},IF({y(c)}=2028,{X('ocx_28')},"
                         f"MAX({X('ocx_floor')},{X('ocx_pct')}*{c}{R['rev']}))))")
    row("ligpay", "Ligado spectrum closing payment", "$mm",
        lambda c, pc, i: f"=IF({y(c)}={X('ligado_year')},{X('ligado_pay')},0)")
    row("jleo", "J-LEO non-dilutive capital (inflow)", "$mm", lambda c, pc, i: f"=IF({y(c)}=2028,{X('jleo')},0)")
    # --- tax
    sec("CASH TAXES (NOL + depreciation)")
    row("capex_tax", "Tax-basis capex placed in service", "$mm", lambda c, pc, i: f"={c}{R['SC']}+{c}{R['ocx']}")
    row("tdep", "Tax depreciation", "$mm",
        lambda c, pc, i: f"=IF({X('bonus_dep')}=1,{c}{R['capex_tax']},SUMIFS($D${R['capex_tax']}:{c}{R['capex_tax']},"
                         f"$D${R['year']}:{c}{R['year']},\">\"&({y(c)}-{life}))/{life}*{c}{R['len']})")
    row("samort", "Spectrum amortisation (15-yr, s.197)", "$mm",
        lambda c, pc, i: f"=IF(AND({y(c)}>={X('ligado_year')},{y(c)}<{X('ligado_year')}+15),{X('ligado_pay')}/15*{c}{R['len']},0)")
    row("ti", "Taxable income", "$mm", lambda c, pc, i: f"={c}{R['ebitda']}-{c}{R['tdep']}-{c}{R['samort']}")
    rnb = rowno[0] + 1
    row("nolb", "NOL pool - beginning", "$mm", lambda c, pc, i: f"={pc}{rnb+2}", first=lambda c: f"={X('nol_open')}")
    row("nolu", "NOL utilised (80% cap)", "$mm",
        lambda c, pc, i: f"=MIN({c}{R['nolb']},MAX(0,{c}{R['ti']})*{X('nol_limit')})")
    row("nole", "NOL pool - ending", "$mm", lambda c, pc, i: f"={c}{R['nolb']}-{c}{R['nolu']}+MAX(0,-{c}{R['ti']})")
    assert R["nole"] == rnb + 2
    row("tax", "Cash taxes", "$mm", lambda c, pc, i: f"={X('tax_rate')}*MAX(0,{c}{R['ti']}-{c}{R['nolu']})", bold=True)
    # --- FCF
    sec("FREE CASH FLOW TO THE FIRM")
    row("ann", "Annualised revenue", "$mm", lambda c, pc, i: f"={c}{R['rev']}/{c}{R['len']}")
    row("dnwc", "Increase in net working capital", "$mm",
        lambda c, pc, i: f"={X('nwc_pct')}*({c}{R['ann']}-{pc}{R['ann']})",
        first=lambda c: f"={X('nwc_pct')}*({c}{R['ann']}-{X('h1_ann_rev')})")
    row("cl", "Contract-liability unwind (revenue already paid in cash)", "$mm",
        lambda c, pc, i: f"=IF(AND({y(c)}>={X('cl_unwind_start')},{y(c)}<{X('cl_unwind_start')}+{X('cl_unwind_yrs')}),"
                         f"{X('contract_liab')}/{X('cl_unwind_yrs')}*{c}{R['len']},0)")
    row("fcf", "FCFF", "$mm",
        lambda c, pc, i: f"={c}{R['ebitda']}-{c}{R['tax']}-{c}{R['netSC']}-{c}{R['ocx']}-{c}{R['ligpay']}+{c}{R['jleo']}"
                         f"-{c}{R['dnwc']}-{c}{R['cl']}", bold=True)
    row("cumfcf", "Cumulative FCFF from 1-Jul-2026", "$mm", lambda c, pc, i: f"=SUM($D${R['fcf']}:{c}{R['fcf']})")
    # --- discounting
    sec("DISCOUNTING (mid-period, from valuation date)")
    row("w", "WACC", "%", lambda c, pc, i: f"=WACC!{c}${WACC_ROW}", fmt="pct2")
    row("comp", "Compounding factor to period end", "f",
        lambda c, pc, i: f"={pc}{R['w']+1}*(1+{c}{R['w']})^{c}{R['len']}",
        first=lambda c: f"=(1+{c}{R['w']})^({c}{R['start']}+{c}{R['len']}-{IREF('val_date')})")
    row("df", "Discount factor (mid-period)", "f",
        lambda c, pc, i: f"=1/({pc}{R['comp']}*(1+{c}{R['w']})^({c}{R['len']}/2))",
        first=lambda c: f"=1/(1+{c}{R['w']})^({c}{R['mid']}-{IREF('val_date')})")
    row("pv", "PV of FCFF", "$mm", lambda c, pc, i: f"={c}{R['fcf']}*{c}{R['df']}", bold=True)
    # terminal column S (display)
    last = PCOLS[-1]
    ws[f"{TCOL}{R['year']}"] = 2041
    ws[f"{TCOL}{R['rev']}"] = f"={last}{R['rev']}*(1+{IREF('g')})"; ws[f"{TCOL}{R['rev']}"].number_format = FMT["$mm"]
    ws[f"{TCOL}{R['ebitda']}"] = f"={last}{R['ebitda']}*(1+{IREF('g')})"; ws[f"{TCOL}{R['ebitda']}"].number_format = FMT["$mm"]
    # --- terminal & valuation (scalars in C)
    sec("TERMINAL VALUE & ENTERPRISE VALUE (col C)")

    def scal(key, label, f, fmt="$mm", bold=False):
        rr = nxt(); R[key] = rr
        ws.cell(rr, 1, label).font = BOLD if bold else BLACK
        c = ws[f"C{rr}"]; c.value = f; c.number_format = FMT.get(fmt, fmt); c.font = BOLD if bold else BLACK
        return rr
    scal("wT", "Terminal WACC", f"={WR['w_t']}", "pct2")
    scal("g", "Terminal growth g", f"={IREF('g')}", "pct2")
    scal("e41", "EBITDA 2041 (normalised)", f"={last}{R['ebitda']}*(1+C{R['g']})")
    scal("maint", "Normalised replacement satellite capex", f"={X('fl_ss')}/{life}/(1-{X('loss_rate')})*{last}{R['U']}*(1+C{R['g']})")
    scal("ocx41", "Other capex 2041", f"={last}{R['ocx']}*(1+C{R['g']})")
    scal("tax41", "Normalised cash tax (tax dep = capex in steady state)",
         f"={X('tax_rate')}*MAX(0,C{R['e41']}-C{R['maint']}-C{R['ocx41']})")
    scal("dnwc41", "Increase in NWC 2041", f"={X('nwc_pct')}*{last}{R['rev']}*C{R['g']}")
    scal("fcf41", "FCFF 2041 (normalised)", f"=C{R['e41']}-C{R['tax41']}-C{R['maint']}-C{R['ocx41']}-C{R['dnwc41']}", bold=True)
    scal("tv", "Terminal value at end-2040 (Gordon, mid-year consistent)",
         f"=C{R['fcf41']}*(1+C{R['wT']})^0.5/(C{R['wT']}-C{R['g']})")
    scal("pvtv", "PV of terminal value", f"=C{R['tv']}/{last}{R['comp']}")
    scal("pvx", "Sum of PV of explicit FCFF", f"=SUM(D{R['pv']}:{last}{R['pv']})")
    scal("ev", "ENTERPRISE VALUE", f"=C{R['pvx']}+C{R['pvtv']}", bold=True)
    scal("tvpct", "Terminal value % of EV", f"=IF(C{R['ev']}<=0,0,C{R['pvtv']}/C{R['ev']})", "%")
    scal("tvmult", "Implied terminal EV / EBITDA (2041)", f"=IF(C{R['e41']}<=0,\"n/m\",C{R['tv']}/C{R['e41']})", "x")
    sec("EQUITY BRIDGE & DILUTION (col C)")
    scal("cash", "+ Cash & liquid assets (pro forma)", f"={CASHREF}")
    scal("debt", "- Non-convertible debt", f"={DEBTREF}")
    scal("conv", "- Convertible notes at face (all treated as debt first)", f"={CONVREF}")
    scal("E0", "Equity value before conversions", f"=C{R['ev']}+C{R['cash']}-C{R['debt']}-C{R['conv']}", bold=True)
    scal("S0", "Basic economic shares (mm)", f"={S0REF}", "mm")
    # dilution engine block
    rr = nxt(); R["dil_hdr"] = rr
    for j, h in enumerate(["Instrument", "Trigger K $", "Shares mm", "D $mm", "Value/sh if prefix converts", "Converts? (1/0)"]):
        ws.cell(rr, 1 + j, h).font = BOLD
    ws.cell(rr, 8, "Order-independent: instrument converts iff value with all K <= its K converting exceeds its K (R-07).").font = GREY
    d0 = rr + 1; d1 = rr + N_INST
    for j in range(N_INST):
        rr = nxt(); ci = INST_FIRST + j
        ws[f"A{rr}"] = f"=Capital_Structure!A{ci}"; ws[f"A{rr}"].font = GREEN
        ws[f"B{rr}"] = f"=Capital_Structure!B{ci}"; ws[f"B{rr}"].font = GREEN; ws[f"B{rr}"].number_format = FMT["$"]
        ws[f"C{rr}"] = f"=Capital_Structure!C{ci}"; ws[f"C{rr}"].font = GREEN; ws[f"C{rr}"].number_format = FMT["f"]
        ws[f"D{rr}"] = f"=Capital_Structure!D{ci}"; ws[f"D{rr}"].font = GREEN; ws[f"D{rr}"].number_format = FMT["$mm"]
        ws[f"E{rr}"] = (f"=($C${R['E0']}+SUMIF($B${d0}:$B${d1},\"<=\"&B{rr},$D${d0}:$D${d1}))/"
                        f"($C${R['S0']}+SUMIF($B${d0}:$B${d1},\"<=\"&B{rr},$C${d0}:$C${d1}))")
        ws[f"E{rr}"].number_format = FMT["$"]
        ws[f"F{rr}"] = f"=IF(E{rr}>B{rr},1,0)"
    scal("Ef", "Equity value after conversions", f"=C{R['E0']}+SUMPRODUCT(F{d0}:F{d1},D{d0}:D{d1})")
    scal("Sf", "Diluted shares (mm)", f"=C{R['S0']}+SUMPRODUCT(F{d0}:F{d1},C{d0}:C{d1})", "mm")
    scal("vps", "VALUE PER SHARE ($, floored at 0 - limited liability)", f"=MAX(0,C{R['Ef']}/C{R['Sf']})", "$", bold=True)
    ws[f"C{R['vps']}"].fill = FILL_Y
    scal("upside", "Upside / (downside) vs current price", f"=C{R['vps']}/{MS['price']}-1", "%")
    R["dil0"], R["dil1"] = d0, d1
    sec("DIAGNOSTICS (col C)")
    scal("fy26", "FY2026 revenue (H1 actual + H2 model)", f"={X('h1_rev')}+D{R['rev']}")
    scal("minfcf", "Peak cumulative funding need (min cumulative FCFF)", f"=MIN(D{R['cumfcf']}:{last}{R['cumfcf']})")
    scal("maxspsat", "Max paying users per satellite (mm)", f"=MAX(D{R['subspersat']}:{last}{R['subspersat']})", "f")
    scal("cagr", "Revenue CAGR 2036-2040", f"=IF(N{R['rev']}<=0,0,({last}{R['rev']}/N{R['rev']})^(1/4)-1)", "%")
    scal("m40", "EBITDA margin 2040", f"={last}{R['margin']}", "%")
    scal("rev30", "Revenue 2030", f"=H{R['rev']}")
    scal("rev35", "Revenue 2035", f"=M{R['rev']}")
    CALC_ROWS[s] = R
    return ws


for s in range(4):
    build_calc(s)

# ================================================================== VALUATION SUMMARY
wsV = wb.create_sheet("Valuation_Summary", 1)
title(wsV, "VALUATION SUMMARY", "Probability-weighted DCF across four scenarios (all live). USD.")
setw(wsV, {"A": 50, "B": 14, "C": 14, "D": 14, "E": 14, "F": 14, "G": 60})
put(wsV, "A4", "Metric", HDR, fill=FILL_H)
for s in range(4):
    put(wsV, f"{CL(2+s)}4", I.SCEN[s], HDR, fill=FILL_H)
put(wsV, "F4", "Prob-weighted", HDR, fill=FILL_H)
metrics = [("Probability", None, "%"), ("Enterprise value ($mm)", "ev", "$mm"), ("PV of terminal value ($mm)", "pvtv", "$mm"),
           ("Terminal value % of EV", "tvpct", "%"), ("Equity before conversions ($mm)", "E0", "$mm"),
           ("Diluted shares (mm)", "Sf", "mm"), ("Value per share ($)", "vps", "$"), ("Upside vs price", "upside", "%"),
           ("Revenue 2030 ($mm)", "rev30", "$mm"), ("Revenue 2035 ($mm)", "rev35", "$mm"),
           ("EBITDA margin 2040", "m40", "%"), ("Implied terminal EV/EBITDA", "tvmult", "x"),
           ("Peak funding need, cumulative FCFF ($mm)", "minfcf", "$mm"), ("Max paying users per satellite (mm)", "maxspsat", "f"),
           ("FY2026 revenue ($mm)", "fy26", "$mm")]
VS = {}
for j, (lab, key, fmt) in enumerate(metrics):
    rr = 5 + j
    put(wsV, f"A{rr}", lab, BOLD if key == "vps" else BLACK)
    for s in range(4):
        f = f"=Inputs!{SCOL[s]}{PROB_ROW}" if key is None else f"=Calc_{I.SCEN[s]}!$C${CALC_ROWS[s][key]}"
        put(wsV, f"{CL(2+s)}{rr}", f, GREEN, fmt)
    if key == "vps" or key is None:   # A3-11: only floored per-share values are probability-weighted
        f = f"=SUM(B{rr}:E{rr})" if key is None else f"=SUMPRODUCT($B$5:$E$5,B{rr}:E{rr})"
        put(wsV, f"F{rr}", f, BOLD, fmt)
    VS[key or "prob"] = rr
PW = f"Valuation_Summary!$F${VS['vps']}"
wsV[f"F{VS['vps']}"].fill = FILL_Y
rr = 5 + len(metrics) + 1
put(wsV, f"A{rr}", "Current share price ($, IBKR)", BOLD); put(wsV, f"F{rr}", f"={MS['price']}", GREEN, "$"); PRICE_ROW = rr; rr += 1
put(wsV, f"A{rr}", "Probability-weighted upside / (downside)", BOLD)
put(wsV, f"F{rr}", f"=F{VS['vps']}/F{PRICE_ROW}-1", BOLD, "%"); rr += 1
put(wsV, f"A{rr}", "Growth-phase cost of equity (Ke)"); put(wsV, f"F{rr}", f"={WR['ke_hi']}", GREEN, "pct2"); KE_ROW = rr; rr += 1
for H in (3, 5):
    put(wsV, f"A{rr}", f"Model-implied expected return p.a. if price converges to value in {H} yrs")
    put(wsV, f"F{rr}", f"=IF(F{VS['vps']}<=0,-1,(F{VS['vps']}*(1+F{KE_ROW})^{H}/F{PRICE_ROW})^(1/{H})-1)", BOLD, "%")
    put(wsV, f"G{rr}", "E[R] = (V x (1+Ke)^H / P)^(1/H) - 1. Replaces the 12% placeholder in the portfolio analysis.", GREY)
    VS[f"er{H}"] = rr; rr += 1
put(wsV, f"A{rr}", "Market-implied EV ($mm)"); put(wsV, f"F{rr}", f"={MKT_EV}", GREEN, "$mm"); rr += 1
put(wsV, f"A{rr}", "Base-case EV / market-implied EV"); put(wsV, f"F{rr}", f"=D{VS['ev']}/{MKT_EV}", BLACK, "x"); rr += 2
ch = BarChart(); ch.type = "col"; ch.title = "Value per share by scenario vs price ($)"; ch.style = 10
ch.add_data(Reference(wsV, min_col=2, max_col=5, min_row=VS['vps'], max_row=VS['vps']), from_rows=True, titles_from_data=False)
ch.set_categories(Reference(wsV, min_col=2, max_col=5, min_row=4, max_row=4))
ch.series[0].tx = SeriesLabel(v="Value per share"); ch.height = 7; ch.width = 16; ch.legend = None
wsV.add_chart(ch, f"A{rr}")

# ================================================================== SENSITIVITY (live grid) + stamped tables
wsS = wb.create_sheet("Sensitivity")
title(wsS, "SENSITIVITY ANALYSIS", "Top grid is LIVE (Base scenario, WACC shift x terminal g, full dilution engine per cell). "
      "Lower tables are generated by engine.py (values stamped; re-run to refresh).")
setw(wsS, {"A": 34, "B": 11, "C": 11, "D": 11, "E": 11, "F": 11, "G": 11, "H": 11})
RB = CALC_ROWS[2]; CB = "Calc_Base"
deltas = [-0.02, -0.01, 0.0, 0.01, 0.02]; gs = [0.015, 0.020, 0.025, 0.030, 0.035]
put(wsS, "A4", "Value per share ($) - Base scenario", BOLD)
put(wsS, "A5", "WACC shift (rows) / terminal g (cols)", GREY)
for j, off in enumerate([-0.01, -0.005, 0.0, 0.005, 0.01]):
    put(wsS, f"{CL(2+j)}5", f"={IREF('g')}+({off})", BLACK, "pct2")   # A3-05: axis centred on live g
for i, dd in enumerate(deltas):
    put(wsS, f"A{6+i}", dd, BLUE, "pct2")
# helper: discount factors per delta (rows 60+)
H0 = 60
put(wsS, f"A{H0-1}", "Helper: discount factors per WACC shift (Base FCFF)", BOLD)
for i in range(5):
    rw = H0 + i * 3
    put(wsS, f"A{rw}", f"=\"WACC+\"&TEXT(A{6+i},\"0.0%\")", BLACK)
    for k, c in enumerate(PCOLS):
        col = CL(2 + k); pcol = CL(1 + k)
        put(wsS, f"{col}{rw}", f"={CB}!{c}{RB['w']}+$A${6+i}", BLACK, "pct2")
        if k == 0:
            put(wsS, f"{col}{rw+1}", f"=(1+{col}{rw})^({CB}!{c}{RB['start']}+{CB}!{c}{RB['len']}-{IREF('val_date')})", BLACK, "f")
            put(wsS, f"{col}{rw+2}", f"=1/(1+{col}{rw})^({CB}!{c}{RB['mid']}-{IREF('val_date')})", BLACK, "f")
        else:
            put(wsS, f"{col}{rw+1}", f"={pcol}{rw+1}*(1+{col}{rw})^{CB}!{c}{RB['len']}", BLACK, "f")
            put(wsS, f"{col}{rw+2}", f"=1/({pcol}{rw+1}*(1+{col}{rw})^({CB}!{c}{RB['len']}/2))", BLACK, "f")
LASTH = CL(1 + len(PCOLS))
# helper: EV / E0 / dilution per grid cell
G0 = 80
put(wsS, f"A{G0-1}", "Helper: per-cell valuation and dilution engine", BOLD)
hdrs = ["Cell", "WACC shift", "g", "EV", "E0"] + [f"conv{j+1}" for j in range(N_INST)] + ["Equity", "Shares", "$/sh"]
for j, h in enumerate(hdrs):
    put(wsS, f"{CL(1+j)}{G0}", h, BOLD)
d0, d1 = f"Capital_Structure!$B${INST_FIRST}", f"Capital_Structure!$B${INST_LAST}"
KR = f"Capital_Structure!$B${INST_FIRST}:$B${INST_LAST}"; NR = f"Capital_Structure!$C${INST_FIRST}:$C${INST_LAST}"
DR = f"Capital_Structure!$D${INST_FIRST}:$D${INST_LAST}"
wT = WR["w_t"]; last = PCOLS[-1]
k = 0
for i in range(5):
    for j in range(5):
        rw = G0 + 1 + k; hr = H0 + i * 3
        put(wsS, f"A{rw}", f"r{i}c{j}")
        put(wsS, f"B{rw}", f"=$A${6+i}", BLACK, "pct2"); put(wsS, f"C{rw}", f"={CL(2+j)}$5", BLACK, "pct2")
        wTd = f"({wT}+B{rw})"
        e41 = f"{CB}!{last}{RB['ebitda']}*(1+C{rw})"
        mnt = f"({CB}!$C${RB['maint']}/(1+{CB}!$C${RB['g']})*(1+C{rw}))"
        ocx = f"{CB}!{last}{RB['ocx']}*(1+C{rw})"
        tax = f"{IREF('tax_rate', 2)}*MAX(0,{e41}-{mnt}-{ocx})"
        fcf41 = f"({e41}-{tax}-{mnt}-{ocx}-{IREF('nwc_pct')}*{CB}!{last}{RB['rev']}*C{rw})"
        put(wsS, f"D{rw}", f"=SUMPRODUCT({CB}!$D${RB['fcf']}:${last}${RB['fcf']},$B${hr+2}:${LASTH}${hr+2})"
                           f"+{fcf41}*(1+{wTd})^0.5/({wTd}-C{rw})/${LASTH}${hr+1}", BLACK, "$mm")
        put(wsS, f"E{rw}", f"=D{rw}+{CASHREF}-{DEBTREF}-{CONVREF}", BLACK, "$mm")
        for m in range(N_INST):
            col = CL(6 + m); Km = f"Capital_Structure!$B${INST_FIRST+m}"
            put(wsS, f"{col}{rw}", f"=IF((E{rw}+SUMIF({KR},\"<=\"&{Km},{DR}))/({S0REF}+SUMIF({KR},\"<=\"&{Km},{NR}))>{Km},1,0)")
        cA, cB = CL(6), CL(5 + N_INST)
        eq, sh, vp = CL(6 + N_INST), CL(7 + N_INST), CL(8 + N_INST)
        eqf = "+".join(f"{CL(6+m)}{rw}*Capital_Structure!$D${INST_FIRST+m}" for m in range(N_INST))
        shf = "+".join(f"{CL(6+m)}{rw}*Capital_Structure!$C${INST_FIRST+m}" for m in range(N_INST))
        put(wsS, f"{eq}{rw}", f"=E{rw}+{eqf}", BLACK, "$mm")
        put(wsS, f"{sh}{rw}", f"={S0REF}+{shf}", BLACK, "mm")
        put(wsS, f"{vp}{rw}", f"=MAX(0,{eq}{rw}/{sh}{rw})", BLACK, "$")
        put(wsS, f"{CL(2+j)}{6+i}", f"={vp}{rw}", BLACK, "$")
        k += 1
wsS["D8"].fill = FILL_Y
put(wsS, "A12", "Check: centre cell equals Calc_Base value per share", GREY)
put(wsS, "D12", f"=D8-{CB}!$C${RB['vps']}", BLACK, "f"); SENS_CHK = "Sensitivity!$D$12"
SENS_STAMP_ROW = 15

# ================================================================== CHECKS
wsK = wb.create_sheet("Checks")
title(wsK, "VALIDATION CHECKS", "Every check is a live formula. Overall status must read ALL PASS before the model is used.")
setw(wsK, {"A": 6, "B": 78, "C": 16, "D": 10, "E": 60})
for j, h in enumerate(["#", "Check", "Value", "Pass?", "Why it matters"]):
    put(wsK, f"{CL(1+j)}4", h, HDR, fill=FILL_H)
chks = []
chks.append(("Scenario probabilities sum to 100%", f"={PROB_SUM}", f"=ABS(C{{r}}-1)<0.0001", "Weights must be a probability distribution."))
chks.append(("Terminal g below risk-free rate", f"={IREF('g')}-{IREF('rf')}", "=C{r}<0", "g > Rf implies firm outgrows economy forever."))
chks.append(("Terminal WACC exceeds g by >= 3pp", f"={WR['w_t']}-{IREF('g')}", "=C{r}>=0.03", "Gordon model explodes as WACC -> g."))
chks.append(("WACC what-if lever is zero", f"={IREF('wacc_shift')}", "=C{r}=0", "Delivered valuation must not carry a manual shift."))
chks.append(("Pro forma cash within 3% of company's 'over $3.7bn'", f"={PF_CASH}", "=AND(C{r}>=3700,C{r}<=3700*1.03)", "Reconciles bridge to company disclosure."))
chks.append(("Debt schedule ties to 10-Q total face $3,022.152mm", f"={DEBT_CHK}", "=ABS(C{r}-3022.152)<0.01", "Catches omitted or double-counted debt."))
chks.append(("Share count ties to 10-Q cover (389.168mm)", f"={S0REF}", "=ABS(C{r}-389.167494)<0.001", "Up-C share classes all included."))
chks.append(("Sensitivity grid centre = Calc_Base value", f"={SENS_CHK}", "=ABS(C{r})<0.005", "Independent re-derivation of Base value."))
for s in range(4):
    Rr = CALC_ROWS[s]; sh = f"Calc_{I.SCEN[s]}"
    chks.append((f"{I.SCEN[s]}: discount factors strictly decreasing", f"=SUMPRODUCT(--({sh}!E{Rr['df']}:R{Rr['df']}>={sh}!D{Rr['df']}:Q{Rr['df']}))",
                 "=C{r}=0", "Detects broken discounting / negative WACC."))
    chks.append((f"{I.SCEN[s]}: fleet never negative", f"=MIN({sh}!D{Rr['end']}:R{Rr['end']})", "=C{r}>=0", "Fleet mechanics."))
    chks.append((f"{I.SCEN[s]}: NOL pool never negative", f"=MIN({sh}!D{Rr['nole']}:R{Rr['nole']})", "=C{r}>=-0.001", "Tax mechanics."))
    chks.append((f"{I.SCEN[s]}: CIP credit used <= pool", f"=SUM({sh}!D{Rr['C']}:R{Rr['C']})-{sh}!C{Rr['pool']}", "=C{r}<=0.001", "No double use of prepaid capex."))
    chks.append((f"{I.SCEN[s]}: EBITDA margin 2040 within 40-75% (or scenario is Distress)",
                 f"={sh}!C{Rr['m40']}", f"=OR(AND(C{{r}}>=0.4,C{{r}}<=0.75),{s}=0)", "Benchmark: Iridium 56.8% (FY25)."))
    chks.append((f"{I.SCEN[s]}: revenue CAGR 2036-40 within g +/- 1.5pp (or Distress)",
                 f"={sh}!C{Rr['cagr']}-{IREF('g')}", f"=OR(ABS(C{{r}})<=0.015,{s}=0)", "Explicit period must be consistent with terminal growth (A3-01)."))
    chks.append((f"{I.SCEN[s]}: paying users per satellite <= 2.5mm",
                 f"={sh}!C{Rr['maxspsat']}", "=C{r}<=2.5", "Capacity plausibility (A2-04)."))
    chks.append((f"{I.SCEN[s]}: converted instruments all have K < value per share",
                 f"=SUMPRODUCT({sh}!F{Rr['dil0']}:F{Rr['dil1']},--({sh}!B{Rr['dil0']}:B{Rr['dil1']}>={sh}!C{Rr['vps']}+0.000001))",
                 "=C{r}=0", "Dilution engine is value-consistent."))
chks.append(("Base FY2026 revenue within guidance ($150-200mm)", f"=Calc_Base!C{CALC_ROWS[2]['fy26']}",
             f"=AND(C{{r}}>={IREF('guide_lo')},C{{r}}<={IREF('guide_hi')})", "Near-term anchor to management guidance."))
chks.append(("Scenario ordering: Distress <= Bear <= Base <= Bull", f"=Valuation_Summary!E{VS['vps']}-Valuation_Summary!B{VS['vps']}",
             f"=AND(Valuation_Summary!B{VS['vps']}<=Valuation_Summary!C{VS['vps']},Valuation_Summary!C{VS['vps']}<=Valuation_Summary!D{VS['vps']},Valuation_Summary!D{VS['vps']}<=Valuation_Summary!E{VS['vps']})",
             "Monotonic scenarios; catches broken scenario wiring."))
chks.append(("Base implied terminal EV/EBITDA within 5-15x", f"=Calc_Base!C{CALC_ROWS[2]['tvmult']}", "=AND(C{r}>=5,C{r}<=15)", "Cross-check Gordon TV vs trading multiples."))
chks.append(("Beta regression uses >= 100 weekly observations", f"={MS['n']}", "=C{r}>=100", "Statistical adequacy."))
first_chk = 5
for j, (lab, val, test, why) in enumerate(chks):
    rr = first_chk + j
    put(wsK, f"A{rr}", j + 1); put(wsK, f"B{rr}", lab); put(wsK, f"C{rr}", val, BLACK, "#,##0.0000")
    put(wsK, f"D{rr}", test.replace("{r}", str(rr)).replace("C{r}", f"C{rr}"), BLACK)
    put(wsK, f"E{rr}", why, GREY)
last_chk = first_chk + len(chks) - 1
put(wsK, "B3", "OVERALL STATUS", BOLD)
put(wsK, "C3", f"=IF(COUNTIF(D{first_chk}:D{last_chk},FALSE)=0,\"ALL PASS\",COUNTIF(D{first_chk}:D{last_chk},FALSE)&\" FAIL\")", BOLD)
wsK["C3"].fill = FILL_Y
CHECK_STATUS = "Checks!$C$3"

# expose layout for other scripts
LAYOUT = dict(IN=IN, CALC_ROWS=CALC_ROWS, VS=VS, MS=MS, WR=WR, PROB_ROW=PROB_ROW, SENS_STAMP_ROW=SENS_STAMP_ROW,
              N_INST=N_INST, INST_FIRST=INST_FIRST, CHECK_STATUS=CHECK_STATUS, PRICE_ROW=PRICE_ROW,
              first_chk=first_chk, last_chk=last_chk)

if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "/home/claude/asts/model_core.xlsx"
    wb.save(out)
    print("saved", out, "checks:", len(chks))
