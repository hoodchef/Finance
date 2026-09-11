"""Agent 3 - independent audit. Recalculated workbook vs engine.py, every scenario, every period."""
import sys
import numpy as np
from openpyxl import load_workbook
sys.path.insert(0, "/home/claude/asts")
import engine as E
import inputs as I
import importlib


def audit(path, verbose=True):
    bx = importlib.import_module("build_xlsx")
    wb = load_workbook(path, data_only=True)
    cols = [chr(ord("D") + i) for i in range(15)]
    rowmap = {  # workbook key -> engine key
        "end": "E", "N": "N", "ret": "R", "G": "G", "avg": "A", "U": "U", "SC": "SC", "C": "C",
        "cov1": "cov1", "cov2": "cov2", "ad1": "ad1", "ad2": "ad2", "subs1": "subs1", "subs2": "subs2",
        "rev1": "rev1", "rev2": "rev2", "gov": "gov", "prod": "prod", "rev": "rev", "cogs": "cogs", "var": "var",
        "lig": "lig", "spec": "spec", "opex": "opex", "netx": "netx", "sbc": "sbc", "ebitda": "ebitda",
        "ocx": "ocx", "ligpay": "ligpay", "jleo": "jleo", "tdep": "tdep", "samort": "samort", "ti": "ti",
        "nole": "nol_e", "tax": "tax", "dnwc": "dnwc", "cl": "cl", "fcf": "fcf", "w": "w", "df": "df",
        "comp": "Ecum", "pv": "pv"}
    scal = {"fcf41": "fcf41", "tv": "tv", "pvtv": "pv_tv", "ev": "ev", "E0": "E0", "Ef": "Ef", "Sf": "Sf", "vps": "vps",
            "maint": "maint", "e41": "e41"}
    worst = 0.0; fails = []
    for s in range(4):
        ws = wb[f"Calc_{I.SCEN[s]}"]; R = bx.CALC_ROWS[s]; r = E.run(s)
        for k, ek in rowmap.items():
            xl = np.array([ws[f"{c}{R[k]}"].value or 0.0 for c in cols], float)
            en = np.asarray(r[ek], float)
            d = np.max(np.abs(xl - en) / np.maximum(1.0, np.abs(en)))
            worst = max(worst, d)
            if d > 1e-6:
                fails.append((I.SCEN[s], k, d, xl[:4], en[:4]))
        for k, ek in scal.items():
            xl = ws[f"C{R[k]}"].value; en = float(r[ek])
            d = abs(xl - en) / max(1.0, abs(en)); worst = max(worst, d)
            if d > 1e-6:
                fails.append((I.SCEN[s], k, d, xl, en))
    wv = wb["Valuation_Summary"]
    pw_xl = wv[f"F{bx.VS['vps']}"].value
    pw_en, _ = E.prob_weighted()
    status = wb["Checks"]["C3"].value
    failed_checks = [(wb["Checks"][f"B{r}"].value, wb["Checks"][f"C{r}"].value)
                     for r in range(bx.LAYOUT["first_chk"], bx.LAYOUT["last_chk"] + 1)
                     if wb["Checks"][f"D{r}"].value is not True]
    if verbose:
        print(f"max relative diff workbook vs engine: {worst:.2e}")
        print(f"prob-weighted: workbook {pw_xl:.4f}  engine {pw_en:.4f}")
        print("Checks sheet:", status)
        for f in fails[:20]:
            print("  DIFF", f)
        for f in failed_checks:
            print("  FAILED CHECK", f)
    return worst, fails, failed_checks, pw_xl


if __name__ == "__main__":
    audit(sys.argv[1] if len(sys.argv) > 1 else "/home/claude/asts/model_core.xlsx")
